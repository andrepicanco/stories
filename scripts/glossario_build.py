"""Converte a planilha do glossário de Cobrança (aba 'Entidades') no JSON lido pela tool `glossario_buscar`.

Determinístico e sem IA: o mesmo Excel sempre gera o mesmo JSON. Uso, na raiz do projeto:

    .\\.venv\\Scripts\\python.exe -m scripts.glossario_build C:\\caminho\\ontologia-minima-cobranca.xlsx

Por padrão o JSON é gravado ao lado da planilha (glossario-cobranca.json). Só entram linhas com status 'ativo'."""
import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from agent.glossary import Glossary
from agent.text import fold

SHEET = "Entidades"
COLUMNS = ("tipo", "nome_canonico", "sinonimos", "nota", "status")
ACTIVE = "ativo"
DEFAULT_NAME = "glossario-cobranca.json"


class BuildError(ValueError):
    pass


def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def _split_synonyms(text: str, canonical: str) -> list[str]:
    """Separa por ';' (o '/' faz parte do nome), sem vazios nem repetidos e sem repetir o nome canônico."""
    seen, out = {fold(canonical)}, []
    for part in text.split(";"):
        synonym = " ".join(part.split())
        if synonym and fold(synonym) not in seen:
            seen.add(fold(synonym))
            out.append(synonym)
    return out


def read_entities(xlsx: Path) -> tuple[list[dict], Counter]:
    """Entradas ativas na ordem da planilha e a contagem das linhas ignoradas por status."""
    from openpyxl import load_workbook   # só este script precisa do openpyxl; o runtime lê JSON

    try:
        wb = load_workbook(xlsx, read_only=True, data_only=True)
    except PermissionError as err:
        raise BuildError("não foi possível abrir a planilha: feche-a no Excel e tente de novo") from err
    except FileNotFoundError as err:
        raise BuildError(f"planilha não encontrada: {xlsx}") from err
    try:
        if SHEET not in wb.sheetnames:
            raise BuildError(f"aba '{SHEET}' não encontrada (abas: {', '.join(wb.sheetnames)})")
        rows = wb[SHEET].iter_rows(values_only=True)
        header = [_cell(h).casefold() for h in (next(rows, None) or ())]
        missing = [c for c in COLUMNS if c not in header]
        if missing:
            raise BuildError(f"aba '{SHEET}' sem a(s) coluna(s): {', '.join(missing)}")
        position = {c: header.index(c) for c in COLUMNS}

        entries, skipped, seen = [], Counter(), {}
        for number, row in enumerate(rows, start=2):
            data = {c: _cell(row[i]) if i < len(row) else "" for c, i in position.items()}
            if not any(data.values()):
                continue
            if fold(data["status"]) != ACTIVE:           # 'revisar', vazio etc. não entram no agente
                skipped[data["status"] or "(sem status)"] += 1
                continue
            if not data["tipo"] or not data["nome_canonico"]:
                raise BuildError(f"linha {number}: 'tipo' e 'nome_canonico' são obrigatórios")
            key = " ".join(fold(data["nome_canonico"]).split())
            if key in seen:
                raise BuildError(f"linha {number}: nome canônico duplicado '{data['nome_canonico']}' "
                                 f"(já aparece na linha {seen[key]})")
            seen[key] = number
            entries.append({
                "tipo": data["tipo"],
                "nome_canonico": data["nome_canonico"],
                "sinonimos": _split_synonyms(data["sinonimos"], data["nome_canonico"]),
                "nota": data["nota"],
                "status": ACTIVE,
            })
        return entries, skipped
    finally:
        wb.close()


def write_glossary(entries: list[dict], out: Path) -> bool:
    """Grava o JSON (sem data/hora, para ser idempotente). Devolve False se o arquivo já estava igual."""
    out = Path(out)
    text = json.dumps(entries, ensure_ascii=False, indent=2) + "\n"
    if out.is_file() and out.read_text(encoding="utf-8") == text:
        return False
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8", newline="\n")
    return True


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("xlsx", type=Path, help="planilha do glossário (aba 'Entidades')")
    parser.add_argument("--out", type=Path, help=f"JSON de saída (padrão: {DEFAULT_NAME} ao lado da planilha)")
    args = parser.parse_args(argv)
    out = args.out or args.xlsx.with_name(DEFAULT_NAME)
    try:
        entries, skipped = read_entities(args.xlsx)
        if not entries:
            raise BuildError("nenhuma entrada com status 'ativo' na planilha")
        changed = write_glossary(entries, out)
    except BuildError as err:
        print(f"Erro: {err}", file=sys.stderr)
        return 1

    glossary = Glossary(entries)
    print(f"{len(entries)} entradas ativas, {glossary.alias_count} aliases -> {out}"
          f"{'' if changed else ' (sem mudanças)'}")
    if skipped:
        print("Ignoradas por status: " + ", ".join(f"{n}x '{s}'" for s, n in sorted(skipped.items())))
    ambiguous = glossary.ambiguous()
    if ambiguous:
        print(f"{len(ambiguous)} alias(es) ambíguo(s): a tool devolve todas as candidatas e o modelo desambigua:")
        for alias, names in sorted(ambiguous.items()):
            print(f"  '{alias}' -> {', '.join(names)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
